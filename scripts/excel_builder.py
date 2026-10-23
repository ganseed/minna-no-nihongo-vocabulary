"""从统一词库生成阅读词库、固定试卷和可重新随机出题的宏版。"""
from __future__ import annotations
import math
import json
import os
import secrets
import sys
import tempfile
import zipfile
from collections import defaultdict
from contextlib import contextmanager
from pathlib import Path

import xlsxwriter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from scripts.random_exam import DIRECTIONS, filter_records, load_dataset, load_vocabulary, select_questions


def project_root() -> Path:
    if not getattr(sys, "frozen", False):
        return Path(__file__).resolve().parents[1]
    executable = Path(sys.executable).resolve()
    if sys.platform == "darwin" and executable.parent.name == "MacOS":
        return executable.parents[3]
    return executable.parent


def resource_file(relative: str) -> Path:
    # 优先保留包外词库的可编辑性；转位或单独移动 app 时回退到内置资源。
    roots = [project_root()]
    if getattr(sys, "frozen", False):
        bundled = getattr(sys, "_MEIPASS", None)
        if bundled:
            roots.append(Path(bundled))
        roots.append(Path(sys.executable).resolve().parent.parent / "Resources")
    return next((root / relative for root in roots if (root / relative).is_file()), roots[0] / relative)


def default_output_dir() -> Path:
    # Keep generated files alongside the app/package, not in a machine-specific
    # absolute Documents path. The package's output/ directory is created on demand.
    return project_root() / "output"


PROJECT_ROOT = project_root()
DATA_FILE = resource_file("data/vocabulary.json")
OUTPUT_DIR = default_output_dir()
MACRO_TEMPLATE = resource_file("template/默写宏模板.xlsm")
THIN = Side(style="thin", color="B8C4CE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="375A7F")
EXAMPLE_FILL = PatternFill("solid", fgColor="D88C4A")
HEADER_FONT = Font(name="SimHei", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="SimHei", size=11, color="172B3A")
JAPANESE_FONT = Font(name="Yu Gothic", size=11, color="172B3A")
ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def load_records(data_path: Path | None = None) -> list[dict]:
    return load_vocabulary(Path(data_path) if data_path else DATA_FILE)


@contextmanager
def atomic_output(path: Path):
    """写完临时文件再替换，生成失败不破坏已有文件。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(prefix=".generating-", suffix=path.suffix, dir=path.parent)
    os.close(handle)
    temporary = Path(temporary)
    try:
        yield temporary
        os.replace(temporary, path)
    except PermissionError as error:
        raise PermissionError(f"无法保存 {path.name}，请关闭正在打开的同名文件，或选择可写的输出文件夹。") from error
    finally:
        temporary.unlink(missing_ok=True)


def style_range(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.font, cell.alignment, cell.border = BODY_FONT, ALIGN, BORDER
            # 词库内容按原文显示，不能把以等号开头的词条解释为公式。
            if isinstance(cell.value, str):
                cell.data_type = "s"


def style_header(ws, row, min_col, max_col, fill=HEADER_FILL):
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=min_col, max_col=max_col):
        for cell in cells:
            cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, fill, ALIGN, BORDER


def prepare_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def row_height(*values: str, minimum=36) -> float:
    """中日文按宽字符估算行高，避免较长释义被截断。"""
    lines = max(sum(max(1, math.ceil(sum(1 if ord(c) > 127 else .55 for c in line) / 27))
                    for line in str(value).split("\n")) for value in values)
    return max(minimum, lines * 15 + 6)


def print_layout(ws, last_row, last_col):
    ws.print_area = f"A1:{last_col}{last_row}"
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"


def build_vocabulary(records, output_dir=None, *, title="大家的日语Ⅰ", unit_label="课程"):
    grouped = defaultdict(list)
    for record in sorted(records, key=lambda r: (r["lesson"], r["order"])):
        grouped[record["lesson"]].append(record)
    if not grouped:
        raise ValueError("所选课程没有词条。")
    wb = prepare_workbook()
    for lesson, items in sorted(grouped.items()):
        prefix = "Lesson" if unit_label == "课程" else unit_label
        ws = wb.create_sheet(f"{prefix}{lesson:02d}")
        ws.append(["日语", "中文"])
        examples = [r for r in items if r["type"] == "example"]
        for record in items:
            if record["type"] != "example":
                ws.append([record["japanese"], record["chinese"]])
        example_row = None
        if examples:
            example_row = ws.max_row + 1
            ws.append(["【例句】", ""])
            for record in examples:
                ws.append([record["japanese"], record["chinese"]])
        style_range(ws, 2, ws.max_row, 1, 2)
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
            row[0].font = JAPANESE_FONT
        style_header(ws, 1, 1, 2)
        if example_row:
            style_header(ws, example_row, 1, 2, EXAMPLE_FILL)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 40
        for row in ws:
            ws.row_dimensions[row[0].row].height = row_height(*(cell.value or "" for cell in row))
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        print_layout(ws, ws.max_row, "B")
    ws = wb.create_sheet("Vocabulary")
    headers = ["id", "lesson", "order", "type", "japanese", "kana", "kanji", "chinese", "remark",
               "part_of_speech", "pitch_accent", "dictionary_form"]
    headers += sorted({key for r in records for key in r} - set(headers))
    ws.append(headers)
    for record in records:
        ws.append([json.dumps(record[key], ensure_ascii=False) if isinstance(record.get(key), (dict, list))
                   else record.get(key, "") for key in headers])
    style_range(ws, 2, ws.max_row, 1, len(headers))
    for column, key in enumerate(headers, 1):
        if key in {"japanese", "kana", "kanji", "dictionary_form"}:
            for row in ws.iter_rows(min_row=2, min_col=column, max_col=column):
                row[0].font = JAPANESE_FONT
    style_header(ws, 1, 1, len(headers))
    for index, width in enumerate([16, 10, 10, 14, 40, 24, 24, 40, 24, 16, 16, 24], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    path = Path(output_dir or OUTPUT_DIR) / f"{title}词库.xlsx"
    with atomic_output(path) as temporary:
        wb.save(temporary)
    return path


def build_exam(records, output_dir=None, *, lessons=None, count=20, direction="zh_to_jp", seed=None, title="大家的日语Ⅰ"):
    """普通版是本次试卷快照，严格按实际题数建表，无空题占位。"""
    questions = select_questions(records, lessons, count, direction, seed)
    wb = prepare_workbook()
    for name, headers in [("默写", ["序号", "题目", "作答", "答案"])]:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for q in questions:
            ws.append([q["number"], q["question"], "", q["answer"]])
            ws.row_dimensions[ws.max_row].height = row_height(q["question"], q["answer"])
        style_range(ws, 2, ws.max_row, 1, 4)
        for row in ws.iter_rows(min_row=2, max_col=4):
            row[3].font = JAPANESE_FONT
            if direction == "jp_to_zh":
                row[1].font = JAPANESE_FONT
        style_header(ws, 1, 1, 4)
        for col, width in {"A": 9, "B": 40, "C": 30, "D": 34}.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "B2"
        ws.sheet_view.showGridLines = False
        print_layout(ws, ws.max_row, "D")
        ws.page_setup.orientation = "landscape"
    settings = wb.create_sheet("本次设置")
    label = next(k for k, v in DIRECTIONS.items() if v == direction)
    for values in [["本次试卷", "内容"], ["方向", label],
                   ["课程", "全部" if lessons is None else ",".join(map(str, lessons))],
                   ["实际题数", len(questions)],
                   ["更换课程或重新出题", "回到桌面工具重新生成；也可使用宏版的“设置”工作表。"],
                   ["打印提示", "题目、作答、答案在同一页横向打印，答案就在作答栏右侧。"]]:
        settings.append(values)
    style_range(settings, 2, settings.max_row, 1, 2)
    style_header(settings, 1, 1, 2)
    settings.column_dimensions["A"].width, settings.column_dimensions["B"].width = 25, 65
    for row in range(2, settings.max_row + 1):
        settings.row_dimensions[row].height = 36
    settings.sheet_view.showGridLines = False
    path = Path(output_dir or OUTPUT_DIR) / f"{title}默写.xlsx"
    with atomic_output(path) as temporary:
        wb.save(temporary)
    return path


def build_macro_exam(records, output_dir=None, *, lessons=None, count=20, direction="zh_to_jp", seed=None, title="大家的日语Ⅰ", unit_label="课程"):
    """保持现有 VBA；用条件边框及动态打印范围适配多题和少题。"""
    if not MACRO_TEMPLATE.is_file():
        raise FileNotFoundError(f"缺少宏模板：{MACRO_TEMPLATE}")
    questions = select_questions(records, lessons, count, direction, seed)
    eligible = [r for r in records if r["chinese"].strip()]
    path = Path(output_dir or OUTPUT_DIR) / f"{title}默写-宏版.xlsm"
    with zipfile.ZipFile(MACRO_TEMPLATE) as archive:
        vba_project = archive.read("xl/vbaProject.bin")
    with atomic_output(path) as temporary, tempfile.TemporaryDirectory() as temp_dir:
        vba_path = Path(temp_dir) / "vbaProject.bin"
        vba_path.write_bytes(vba_project)
        with xlsxwriter.Workbook(temporary, {"strings_to_formulas": False, "strings_to_urls": False}) as wb:
            wb.add_vba_project(vba_path)
            basic = {"font_name": "SimHei", "font_size": 11, "font_color": "#172B3A",
                     "align": "center", "valign": "vcenter", "text_wrap": True}
            body = wb.add_format(basic)
            japanese_body = wb.add_format({**basic, "font_name": "Yu Gothic"})
            header = wb.add_format({**basic, "bold": True, "font_color": "#FFFFFF",
                                    "bg_color": "#375A7F", "border": 1})
            borders = wb.add_format({"border": 1, "border_color": "#B8C4CE"})
            settings = wb.add_worksheet("设置")
            source = wb.add_worksheet("题库")
            exam = wb.add_worksheet("默写")
            for sheet in (settings, source, exam):
                sheet.hide_gridlines(2)
            settings.set_column("A:A", 24, body)
            settings.set_column("B:B", 60, body)
            settings.set_default_row(32)
            label = next(k for k, v in DIRECTIONS.items() if v == direction)
            settings_rows = [["默写设置", "选择值"], ["方向", label],
                [unit_label, "全部" if lessons is None else ",".join(map(str, lessons))],
                ["随机数量", "全部" if count is None else count],
                ["使用方式", "修改设置后点击右侧“随机生成”；查看题目请切换到“默写”。"],
                ["打印提示", "题目、作答、答案在同一页横向打印，范围随实际题数变化。"],
                ["课程填写示例", "单课：3；多课：1,3,5；全部：全部（此处不支持范围写法）"],
                ["实际题数", ""], ["题库范围", f"保留全部 {len(eligible)} 条可用题目，可继续更换课程。"]]
            for row, values in enumerate(settings_rows):
                settings.write_row(row, 0, values, header if row == 0 else body)
            settings.write_formula("B8", "=COUNTA('默写'!A:A)-1", body, len(questions))
            settings.data_validation("B2", {"validate": "list", "source": list(DIRECTIONS), "error_type": "stop", "error_message": "请选择中译日或日译中。", "ignore_blank": False})
            settings.data_validation("B4", {"validate": "custom", "value": '=IF(B4="全部",TRUE,IF(ISNUMBER(B4),AND(B4=INT(B4),B4>0,B4<=1048575),FALSE))', "error_type": "stop", "error_message": "请输入正整数或“全部”。", "ignore_blank": False})
            settings.insert_button("D2", {"macro": "RegenerateExam", "caption": "随机生成", "width": 150, "height": 42})
            settings.merge_range("D5:F6", "每次点击都会重新随机，无需回到软件生成。", body)
            settings.write_row("H1", [f"可选{unit_label}编号", "可用题数"], header)
            settings.write_row("H2", ["全部", len(eligible)], body)
            for row, unit in enumerate(sorted({r["lesson"] for r in eligible}), 2):
                settings.write_row(row, 7, [unit, sum(r["lesson"] == unit for r in eligible)], body)
            settings.set_column("H:I", 18, body)
            settings.data_validation("B3", {"validate": "list", "source": f"=$H$2:$H${len({r['lesson'] for r in eligible})+2}", "show_error": False})
            source.write_row(0, 0, ["lesson", "japanese", "chinese", "type", "id"], header)
            for row, record in enumerate(eligible, 1):
                source.write_row(row, 0, [record[k] for k in ["lesson", "japanese", "chinese", "type", "id"]], body)
                source.write(row, 1, record["japanese"], japanese_body)
            source.hide()
            for sheet in (exam,):
                sheet.set_column("A:A", 9, body)
                sheet.set_column("B:B", 40, japanese_body if direction == "jp_to_zh" else body)
                sheet.set_column("C:C", 30, body)
                sheet.set_default_row(36)
                sheet.freeze_panes(1, 1)
                sheet.repeat_rows(0)
                sheet.set_paper(9)
                sheet.set_landscape()
                sheet.fit_to_pages(1, 0)
                sheet.set_footer("&C第 &P 页 / 共 &N 页")
                # 条件格式覆盖整个可用题数，但不创建占位单元格。
                # VBA 清除旧题后，旧边框随 A 列变空自动消失；增题时自动补齐。
                sheet.conditional_format(1, 0, len(eligible), 3,
                    {"type": "formula", "criteria": '$A2<>""', "format": borders})
            exam.set_column("D:D", 34, japanese_body)
            exam.set_column("F:F", 28, body)
            exam.insert_button("F2", {"macro": "RegenerateExam", "caption": "重新随机生成", "width": 150, "height": 42})
            exam.write_url("F5", "internal:'设置'!B2", body, "修改课程、方向和题数")
            exam.write("F7", "每次点击重新抽题并清空旧作答；答案在右侧D列。", body)
            exam.write_row(0, 0, ["序号", "题目", "作答", "答案"], header)
            for row, q in enumerate(questions, 1):
                exam.write(row, 0, row, body)
                exam.write(row, 1, q["question"], japanese_body if direction == "jp_to_zh" else body)
                exam.write_blank(row, 2, None, body)
                exam.write(row, 3, q["answer"], japanese_body)
            # Excel 原生动态打印区域：题数由 VBA 写入的 A 列决定，不受历史最多题数影响。
            for name in ("默写",):
                wb.define_name(f"'{name}'!_xlnm.Print_Area",
                    f"='{name}'!$A$1:INDEX('{name}'!$D:$D,MAX(2,COUNTA('默写'!$A:$A)))")
            settings.activate()
    return path


def build_workbooks(data_path=None, output_dir=None, *, lessons=None, count=20, direction="zh_to_jp", seed=None):
    records, metadata = load_dataset(Path(data_path) if data_path else DATA_FILE)
    title, unit_label = metadata["title"], metadata["unit_label"]
    selected = filter_records(records, lessons)
    # 一次生成中的普通版和宏版初始试卷保持一致；下次点击仍重新随机。
    seed = secrets.randbits(64) if seed is None else seed
    # 出题设置错误时，在写任何文件前立即提示。
    select_questions(records, lessons, count, direction, seed)
    destination = Path(output_dir or OUTPUT_DIR)
    destination.mkdir(parents=True, exist_ok=True)
    # 先完成全部生成，再替换正式输出，避免中途生成失败留下半套新文件。
    with tempfile.TemporaryDirectory(prefix=".generation-", dir=destination) as stage:
        stage = Path(stage)
        paths = (build_vocabulary(selected, stage, title=title, unit_label=unit_label),
                 build_exam(records, stage, lessons=lessons, count=count, direction=direction, seed=seed, title=title),
                 build_macro_exam(records, stage, lessons=lessons, count=count, direction=direction, seed=seed, title=title, unit_label=unit_label))
        result = []
        for source in paths:
            target = destination / source.name
            with atomic_output(target) as temporary:
                temporary.write_bytes(source.read_bytes())
            result.append(target)
    return tuple(result)


if __name__ == "__main__":
    build_workbooks()
