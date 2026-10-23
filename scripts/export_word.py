"""将统一词库导出为 Word 文档。"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.random_exam import filter_records, load_dataset
from scripts.excel_builder import atomic_output


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def export_word(data_path: Path, output_path: Path, *, lessons=None) -> Path:
    """生成每课分节、例句置后的可编辑 Word 文档。"""
    all_records, metadata = load_dataset(data_path)
    records = filter_records(all_records, lessons)
    if not records:
        raise ValueError("所选课程没有词条。")
    grouped: dict[int, list[dict]] = defaultdict(list)
    for record in sorted(records, key=lambda r: (r["lesson"], r["order"])):
        grouped[record["lesson"]].append(record)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "SimHei"
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimHei")
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].font.name = "SimHei"
        document.styles[name].element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimHei")
    normal.font.size = Pt(10.5)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    for index, (lesson, items) in enumerate(sorted(grouped.items())):
        if index:
            document.add_page_break()
        unit_label = "Lesson" if metadata["unit_label"] == "课程" else metadata["unit_label"]
        heading = document.add_heading(f"{unit_label} {lesson:02d}", level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        regular = [item for item in items if item["type"] != "example"]
        examples = [item for item in items if item["type"] == "example"]
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width, table.columns[1].width = Cm(9), Cm(8.4)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "日语", "中文"
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for item in regular:
            cells = table.add_row().cells
            cells[0].width, cells[1].width = Cm(9), Cm(8.4)
            cells[0].text, cells[1].text = item["japanese"], item["chinese"]
            for run in cells[0].paragraphs[0].runs:
                run.font.name = "Yu Gothic"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Yu Gothic")
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if examples:
            document.add_heading("【例句】", level=2)
            for item in examples:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(item["japanese"] + "\n")
                run.font.name = "Yu Gothic"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Yu Gothic")
                paragraph.add_run(item["chinese"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with atomic_output(output_path) as temporary:
        document.save(temporary)
    return output_path


if __name__ == "__main__":
    export_word(PROJECT_ROOT / "data" / "vocabulary.json", PROJECT_ROOT / "output" / "大家的日语Ⅰ词库.docx")
