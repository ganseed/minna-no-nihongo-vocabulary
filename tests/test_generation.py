"""真实生成回归：空行、扩题、筛选、非法数据和导出一致性。"""
from __future__ import annotations
import copy
import contextlib
import io
import json
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook
from scripts.excel_builder import (MACRO_TEMPLATE, atomic_output, build_exam, build_macro_exam,
                                   build_vocabulary, build_workbooks)
from scripts.export_word import export_word
from scripts.random_exam import load_vocabulary, parse_count, parse_lessons, select_questions

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
ROOT = Path(__file__).resolve().parents[1]


def records(count=42, lesson=3):
    return [{"id": f"L{lesson:02}-{n:03}", "lesson": lesson, "order": n,
             "type": "word", "japanese": f"かな{n}\n漢字{n}", "kana": f"かな{n}",
             "kanji": f"漢字{n}", "chinese": f"中文{n}", "remark": ""}
            for n in range(1, count + 1)]


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_chinese_uses_simhei_while_japanese_and_answers_keep_font(self):
        for builder in (build_exam, build_macro_exam):
            for direction in ("zh_to_jp", "jp_to_zh"):
                path = builder(records(), self.root, count=2, direction=direction)
                wb = load_workbook(path)
                ws = wb["默写"]
                self.assertEqual(ws["A1"].font.name, "SimHei")
                self.assertEqual(ws["A2"].font.name, "SimHei")
                self.assertEqual(ws["C2"].font.name, "SimHei")
                self.assertEqual(ws["B2"].font.name, "Yu Gothic" if direction == "jp_to_zh" else "SimHei")
                self.assertEqual(ws["D2"].font.name, "Yu Gothic")
                if builder == build_macro_exam:
                    self.assertEqual(ws.column_dimensions["D"].font.name, "Yu Gothic")
                wb.close()
        wb = load_workbook(build_vocabulary(records(), self.root))
        self.assertEqual(wb["Lesson03"]["A2"].font.name, "Yu Gothic")
        self.assertEqual(wb["Lesson03"]["B2"].font.name, "SimHei")
        wb.close()

    def test_twenty_questions_means_twenty_rows_and_inline_answers(self):
        path = build_exam(records(250), self.root, count=20, seed=1)
        wb = load_workbook(path, data_only=True)
        self.assertEqual(wb["默写"].max_row, 21)
        self.assertNotIn("答案", wb.sheetnames)
        self.assertTrue(wb["默写"]["D2"].value)
        self.assertEqual(str(wb["默写"].print_area), "'默写'!$A$1:$D$21")
        self.assertEqual(wb["默写"].max_column, 4)
        self.assertTrue(all(c.value is None for row in wb["默写"].iter_rows(min_row=2, min_col=3, max_col=3) for c in row))

    def test_lesson_filter_and_shortfall_do_not_create_empty_questions(self):
        pool = records(7, 3) + records(15, 8)
        path = build_exam(pool, self.root, lessons=[3], count=200)
        wb = load_workbook(path, data_only=True)
        self.assertEqual(wb["默写"].max_row, 8)
        self.assertEqual(wb["本次设置"]["B4"].value, 7)

    def test_expand_and_shrink_normal_exam(self):
        for count in [20, 200, 50, 250, 20]:
            wb = load_workbook(build_exam(records(250), self.root, count=count))
            self.assertEqual(wb["默写"].max_row, count + 1)
            self.assertEqual(wb["默写"][f"C{count + 1}"].border.bottom.style, "thin")
            self.assertEqual(str(wb["默写"].print_area), f"'默写'!$A$1:$D${count+1}")

    def test_macro_has_adaptive_borders_without_preformatted_blank_rows(self):
        for count in [20, 200, 20]:
            path = build_macro_exam(records(250), self.root, count=count, seed=1)
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                wb = load_workbook(path, data_only=True)
            exam = wb["默写"]
            self.assertEqual(max(c.row for row in exam.iter_rows(min_col=1, max_col=1) for c in row if c.value is not None), count + 1)
            self.assertIsNone(exam["A252"].value)
            self.assertIsNone(exam["A252"].border.bottom.style)
            self.assertFalse(exam.column_dimensions["D"].hidden)
            self.assertTrue(exam["D2"].value)
            self.assertNotIn("答案", wb.sheetnames)
            self.assertEqual(wb["题库"].sheet_state, "hidden")
            with zipfile.ZipFile(path) as z:
                xml = ET.fromstring(z.read("xl/worksheets/sheet3.xml"))
                rules = xml.findall("s:conditionalFormatting", NS)
                self.assertEqual(rules[0].get("sqref"), "A2:D251")
                self.assertEqual(rules[0].find("s:cfRule/s:formula", NS).text, '$A2<>""')
                names = ET.fromstring(z.read("xl/workbook.xml")).findall("s:definedNames/s:definedName", NS)
                print_names = [n for n in names if n.get("name") == "_xlnm.Print_Area"]
                self.assertEqual(len(print_names), 1)
                self.assertIn("$D:$D", print_names[0].text)
                self.assertTrue(all("COUNTA('默写'!$A:$A)" in n.text for n in print_names))
                with zipfile.ZipFile(MACRO_TEMPLATE) as template:
                    self.assertEqual(z.read("xl/vbaProject.bin"), template.read("xl/vbaProject.bin"))
                # 动态边框确实包含四边，不只检查条件公式的存在。
                styles = ET.fromstring(z.read("xl/styles.xml"))
                border = styles.find("s:dxfs/s:dxf/s:border", NS)
                self.assertTrue(all(border.find(f"s:{edge}", NS).get("style") == "thin"
                                    for edge in ["left", "right", "top", "bottom"]))

    def test_each_generation_draws_fresh_questions_and_macro_buttons_remain(self):
        from unittest.mock import patch
        path = self.root / "input.json"
        path.write_text(json.dumps(records(100)), encoding="utf-8")
        draws = []
        with patch("scripts.excel_builder.secrets.randbits", side_effect=[101, 202]) as entropy:
            for _ in range(2):
                paths = build_workbooks(path, self.root)
                wb = load_workbook(paths[1], data_only=True)
                draws.append([wb["默写"][f"B{r}"].value for r in range(2, 22)])
            self.assertEqual(entropy.call_count, 2)
        self.assertNotEqual(*draws)
        with zipfile.ZipFile(paths[2]) as z:
            drawings = [z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".vml")]
            self.assertEqual(sum(t.count("<x:FmlaMacro>") for t in drawings), 2)
            self.assertTrue(all("RegenerateExam" in t for t in drawings))
            with zipfile.ZipFile(MACRO_TEMPLATE) as original:
                self.assertEqual(z.read("xl/vbaProject.bin"), original.read("xl/vbaProject.bin"))

    def test_custom_library_and_future_nonconsecutive_groups(self):
        pool = records(4, 51) + records(3, 108)
        pool[0]["tags"] = ["旅行", "车站"]
        path = self.root / "custom.json"
        path.write_text(json.dumps({"title": "旅行与工作", "unit_label": "主题", "records": pool}), encoding="utf-8")
        paths = build_workbooks(path, self.root, lessons=[108], count=30)
        self.assertEqual(paths[0].name, "旅行与工作词库.xlsx")
        self.assertEqual(paths[2].name, "旅行与工作默写-宏版.xlsm")
        vocab = load_workbook(paths[0])
        self.assertIn("主题108", vocab.sheetnames)
        exam = load_workbook(paths[1], data_only=True)["默写"]
        self.assertEqual(exam.max_row, 4)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            macro = load_workbook(paths[2], data_only=True)
        self.assertEqual(macro["题库"].max_row, 8)
        self.assertEqual([macro["设置"][f"H{r}"].value for r in range(2,5)], ["全部",51,108])
        self.assertEqual(macro["设置"]["A3"].value, "主题")
        self.assertEqual(macro["设置"]["B3"].value, "108")
        extra = load_workbook(build_vocabulary(pool, self.root))["Vocabulary"]
        self.assertIn("tags", [c.value for c in extra[1]])

    def test_minimal_custom_library_without_courses_still_works(self):
        path = self.root / "自建词库.json"
        path.write_text(json.dumps([{"id":"station", "japanese":"えき", "chinese":"车站"}]), encoding="utf-8")
        loaded = load_vocabulary(path)
        self.assertEqual(loaded[0]["lesson"], 1)
        paths = build_workbooks(path, self.root, count=None)
        self.assertEqual(paths[0].name, "自建词库词库.xlsx")
        self.assertEqual(load_workbook(paths[1])["默写"].max_row, 2)

    def test_all_actual_lessons_only_no_empty_lesson_tabs(self):
        wb = load_workbook(build_vocabulary(records(2, 3) + records(1, 8), self.root))
        self.assertEqual(wb.sheetnames, ["Lesson03", "Lesson08", "Vocabulary"])

    def test_selected_word_export_has_no_leading_blank_page(self):
        path = self.root / "test.json"
        path.write_text(json.dumps(records(2, 3) + records(3, 8)), encoding="utf-8")
        doc = Document(export_word(path, self.root / "words.docx", lessons=[8]))
        self.assertEqual(doc.paragraphs[0].text, "Lesson 08")
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.tables[0].rows), 4)
        self.assertNotIn('w:type="page"', doc.paragraphs[0]._p.xml)

    def test_formula_like_vocabulary_is_literal(self):
        pool = records(1)
        pool[0]["chinese"] = "=1+1"
        normal = load_workbook(build_exam(pool, self.root))
        self.assertEqual(normal["默写"]["B2"].data_type, "s")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            macro = load_workbook(build_macro_exam(pool, self.root))
        self.assertEqual(macro["默写"]["B2"].data_type, "s")

    def test_atomic_output_preserves_previous_file_after_failure(self):
        path = self.root / "existing.xlsx"
        path.write_bytes(b"previous file")
        with self.assertRaises(RuntimeError):
            with atomic_output(path) as temporary:
                temporary.write_bytes(b"incomplete")
                raise RuntimeError("export failed")
        self.assertEqual(path.read_bytes(), b"previous file")
        self.assertEqual(list(self.root.glob(".generating-*")), [])

    def test_generation_does_not_depend_on_chinese_console_support(self):
        path = self.root / "input.json"
        path.write_text(json.dumps(records(2)), encoding="utf-8")
        with io.TextIOWrapper(io.BytesIO(), encoding="cp1252") as english_console:
            with contextlib.redirect_stdout(english_console):
                paths = build_workbooks(path, self.root)
        self.assertTrue(all(p.is_file() for p in paths))

    def test_real_library_single_lesson_full_generation(self):
        paths = build_workbooks(ROOT / "data/vocabulary.json", self.root, lessons=[3], count=200, seed=2)
        selected = [r for r in load_vocabulary(ROOT / "data/vocabulary.json") if r["lesson"] == 3 and r["chinese"].strip()]
        wb = load_workbook(paths[1], data_only=True)
        self.assertEqual(wb["默写"].max_row, len(selected) + 1)
        self.assertNotIn("答案", wb.sheetnames)
        self.assertTrue(wb["默写"]["D2"].value)

    def test_normal_and_macro_initial_questions_match_without_fixed_seed(self):
        path = self.root / "input.json"
        path.write_text(json.dumps(records(100)), encoding="utf-8")
        paths = build_workbooks(path, self.root)
        normal = load_workbook(paths[1], data_only=True)["默写"]
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            macro = load_workbook(paths[2], data_only=True)["默写"]
        self.assertEqual([normal[f"B{r}"].value for r in range(2, 22)],
                         [macro[f"B{r}"].value for r in range(2, 22)])


class SelectionTests(unittest.TestCase):
    def test_courses_chinese_delimiters_ranges_and_deduplication(self):
        self.assertEqual(parse_lessons("1-3，3、5", range(1, 51)), [1, 2, 3, 5])
        self.assertIsNone(parse_lessons("全部", [1, 2]))

    def test_invalid_courses(self):
        for value in ["0", "51", "3-1", "1,,3", "a", "2"]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_lessons(value, [1, 3])

    def test_invalid_counts(self):
        for value in [0, -1, True, "0", "-2", "2.5", "", "a"]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_count(value)
        self.assertIsNone(parse_count("全部"))
        self.assertEqual(parse_count(" 30 "), 30)

    def test_seed_no_duplicates_no_input_mutation_and_directions(self):
        pool = records(12)
        before = copy.deepcopy(pool)
        a = select_questions(pool, count=7, seed=7)
        self.assertEqual(a, select_questions(pool, count=7, seed=7))
        self.assertEqual(len({q["id"] for q in a}), 7)
        self.assertEqual(pool, before)
        self.assertEqual(a[0]["question"], a[0]["chinese"])
        b = select_questions(pool, count=1, direction="jp_to_zh")
        self.assertEqual(b[0]["question"], b[0]["japanese"])

    def test_all_and_blank_translations(self):
        pool = records(3)
        pool[0]["chinese"] = "  "
        self.assertEqual(len(select_questions(pool, count=None)), 2)
        with self.assertRaises(ValueError):
            select_questions(pool, lessons=[])
        with self.assertRaises(ValueError):
            select_questions([], direction="wrong")

    def test_validation_reports_malformed_and_duplicate_records(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "data.json"
            for value in [{}, [], [1], [{"id": "missing"}], records(1) * 2]:
                path.write_text(json.dumps(value), encoding="utf-8")
                with self.subTest(value=value), self.assertRaises(ValueError):
                    load_vocabulary(path)
            path.write_text(json.dumps(records(1)), encoding="utf-8-sig")
            self.assertEqual(len(load_vocabulary(path)), 1)


if __name__ == "__main__":
    unittest.main()
